"""Assemble the cached study pack from source materials.

Reads the 4 PPTX decks, the term exam review PDF, and the practice exam DOCX
from c:\\Users\\gardn\\Downloads, normalizes everything to ASCII, and writes
study_pack.txt to the project root in a deterministic order so the prompt
cache hashes consistently across runs.

Run: python -m host.build_study_pack
"""
from __future__ import annotations

import unicodedata
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation


PROJECT_ROOT = Path(__file__).resolve().parent.parent
SOURCES_DIR = Path("c:/Users/gardn/Downloads")
OUTPUT_FILE = PROJECT_ROOT / "study_pack.txt"

# Ordered list -- alphabetical by filename -- so the cached prefix is stable.
DECKS = [
    "2 - RiskReward, Stock Valuation (2) (1).pptx",
    "CFA_2024_L1_Masterclass_Equity (3).pptx",
    "EMH - Market Efficiency.pptx",
    "Security_Market_Indexes.pptx",
]
REVIEW_PDF = "Term Exam 3 - Review - Final.pdf"
PRACTICE_DOCX = "Practice Exam - Final (Equities) - noKEY.docx"


# Hand-written formula appendix -- this is the load-bearing reference Claude
# uses to answer computational questions. Order and wording are stable.
FORMULA_APPENDIX = """=== FORMULA APPENDIX ===

MARKET STRUCTURE / MARGIN
-------------------------
Margin call price (long position):
  P_call = P0 * (1 - IM) / (1 - MM)
  where IM = initial margin %, MM = maintenance margin % (decimals).

Leverage ratio = 1 / IM
  IM 50% -> 2x; IM 25% -> 4x; IM 40% -> 2.5x.

Return on margin (long, no dividends, no costs):
  R = (P_end - P0) / (IM * P0)
  Equivalent: R = HPR / IM, where HPR = (P_end - P0) / P0.

With dividends D and borrowing cost C and transaction cost T (if STATED):
  R = ((P_end - P0) + D - C - T) / (IM * P0)

Maintenance check (will I get a margin call?):
  Equity / Market value < MM%  -> margin call.
  Equity = Current value - Borrowed amount; Borrowed = P0 * (1 - IM) * shares.

VALUATION
---------
Single-stage Gordon DDM (constant growth):
  V0 = D1 / (k - g)
  Requires: k > g, constant g forever.

Preferred stock value = D / k
  (Special case of DDM with g = 0.)

Multi-stage DDM (two-stage example, explicit dividends D1, D2, then constant g):
  TV2 = D2 * (1 + g) / (k - g)         (terminal value at end of year 2)
  V0 = D1 / (1 + k) + D2 / (1 + k)^2 + TV2 / (1 + k)^2

Justified P/E (leading): P0/E1 = (1 - b) / (k - g)
  Shorthand: Justified P/E = (1 - b) / (k - g)
  where b = retention ratio, (1 - b) = payout ratio.

Sustainable growth: g = b * ROE
  Drives the link between retention, ROE, and stock-price growth.

Required rate of return (CAPM): k = Rf + Beta * (Rm - Rf)
P/S ratio = Price per share / Sales per share
P/B ratio = Price per share / Book value per share
ROE = Net income / Average book value of equity
ROA = Net income / Total assets

INCOME STATEMENT FORECASTING
----------------------------
COGS = Revenue * (1 - Gross margin)
Degree of operating leverage (DOL) = % change in operating income / % change in sales

INDEXES
-------
Price-weighted index: sum of prices / divisor (divisor adjusts for splits).
Value/market-cap weighted: sum of (price * shares) / base.
Equal-weighted: average of % returns.
Total return = price return + dividend return.

EMH
---
Weak form: prices reflect past prices. Technical analysis fails.
Semi-strong: prices reflect all public info. Fundamental analysis fails on average.
Strong form: prices reflect all info, including private/insider. Insider trading fails.
Anomalies (evidence against EMH): low P/E premium, small-cap premium, January effect.

EQUITY OVERVIEW
---------------
Common stock: residual claim, voting rights, last paid in liquidation.
Preferred stock: fixed dividend (cumulative or non), no voting (typically), priority over common.
Cost of equity > cost of preferred > cost of debt (after-tax).
Statutory voting: one share = one vote per director seat.
Cumulative voting: shares * seats votes, distributable across candidates (better for minority shareholders).

"""


def _to_ascii(text: str) -> str:
    """Normalize Unicode and drop characters that don't survive ASCII encoding."""
    text = unicodedata.normalize("NFKD", text)
    replacements = {
        "—": "-", "–": "-",
        "‘": "'", "’": "'",
        "“": '"', "”": '"',
        "…": "...", "°": " deg",
        "•": "-", "×": "x", "÷": "/",
        "≤": "<=", "≥": ">=", "≠": "!=",
        "α": "alpha", "β": "beta", "σ": "sigma",
        "μ": "mu", "√": "sqrt", "∞": "inf",
        "π": "pi", "∑": "sum",
    }
    for u, a in replacements.items():
        text = text.replace(u, a)
    return text.encode("ascii", errors="ignore").decode("ascii")


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return _to_ascii("\n".join(pages))


def extract_pptx_text(path: Path) -> str:
    prs = Presentation(str(path))
    chunks: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks.append(f"\n--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = "".join(run.text for run in para.runs)
                    if txt.strip():
                        chunks.append(txt)
    return _to_ascii("\n".join(chunks))


def extract_docx_text(path: Path) -> str:
    doc = Document(str(path))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    return _to_ascii("\n".join(paras))


def assemble() -> str:
    parts: list[str] = [FORMULA_APPENDIX]

    parts.append("=== TERM EXAM 3 REVIEW ===\n")
    parts.append(extract_pdf_text(SOURCES_DIR / REVIEW_PDF))

    for deck in DECKS:
        parts.append(f"\n=== DECK: {deck} ===\n")
        parts.append(extract_pptx_text(SOURCES_DIR / deck))

    parts.append("\n=== PRACTICE EXAM ===\n")
    parts.append(extract_docx_text(SOURCES_DIR / PRACTICE_DOCX))

    return "\n".join(parts)


def main() -> None:
    pack = assemble()
    OUTPUT_FILE.write_text(pack, encoding="ascii")
    size_k = len(pack) // 1000
    print(f"Wrote {OUTPUT_FILE} ({size_k}K chars, ~{size_k // 4}K tokens estimated)")


if __name__ == "__main__":
    main()
